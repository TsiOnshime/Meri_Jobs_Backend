"""DOCX text extraction (python-docx)."""
"""DOCX text extraction (python-docx)."""
from .common import extract_fields


def extract_raw_text(file_path: str) -> str:
    import docx  # python-docx

    document = docx.Document(file_path)
    chunks = [p.text for p in document.paragraphs if p.text.strip()]

    # Multi-column CVs often put content in tables, not plain paragraphs.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    chunks.append(cell.text.strip())

    return "\n".join(chunks).strip()


def parse(file_path: str) -> dict:
    raw_text = extract_raw_text(file_path)
    return extract_fields(raw_text)